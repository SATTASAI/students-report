# สร้าง public/js/growth-doh619.js จากไฟล์โปรแกรมกชกร (ชีต age / ht)
# python3 dev/build-doh619.py kotchakorn_v2_20210404.xlsx public/js/growth-doh619.js
import openpyxl, json, sys
wb=openpyxl.load_workbook(sys.argv[1])
SEX={'1':'ช','2':'ญ'}
def num(v):
    v=float(v); r=round(v,4)
    return int(r) if r==int(r) else r
age=[r for r in wb['age'].iter_rows(min_row=2,values_only=True) if r[0] is not None]
ht=[r for r in wb['ht'].iter_rows(min_row=2,values_only=True) if r[0] is not None]
out={'wfh':{},'hfa':{},'wfa':{}}
for s,name in SEX.items():
    # อายุ 60 เดือน (WHO) + 61–71 (ค่าคำนวณเชื่อมช่วง) + 72–228 (กรมอนามัย 2564)
    rows={int(str(r[0])[1:]):r for r in age if str(r[0])[0]==s and (r[-1] in (2,3) or str(r[0])[1:]=='60')}
    ms=sorted(rows); assert ms==list(range(60,229)),(ms[0],ms[-1],len(ms))
    # age cols: code, med_wt, sd_wt, -2, -1.5, +1.5, +2, med_ht, sd_ht, -2, -1.5, +1.5, +2, cat
    out['wfa'][name]={'from':60,'step':1,'rows':[[num(rows[m][i]) for i in (3,4,1,5,6)] for m in ms]}
    out['hfa'][name]={'from':60,'step':1,'rows':[[num(rows[m][i]) for i in (9,10,7,11,12)] for m in ms]}
    # ht cols: code, median, sd, -2, -1.5, +1.5, +2, +3, cat — ส่วนสูง 120 ซม. ขึ้นไป (กรมอนามัย)
    hr={int(float(str(r[0])[1:])):r for r in ht if str(r[0])[0]==s and r[-1]==2}
    hs=sorted(hr); assert hs==list(range(120,hs[-1]+1))
    out['wfh'][name]={'from':120,'step':1,'rows':[[num(hr[h][i]) for i in (3,4,1,5,6,7)] for h in hs]}
for k in out:
    for s in out[k]:
        for row in out[k][s]['rows']: assert row==sorted(row),(k,s,row)
def tab(t): return '{ from: %d, step: %d, rows: [\n%s\n  ] }'%(t['from'],t['step'],',\n'.join('    '+json.dumps(r,separators=(',',':')) for r in t['rows']))
js=['''// เกณฑ์อ้างอิงการเจริญเติบโตของเด็กอายุ 6–19 ปี กรมอนามัย (เผยแพร่ พ.ศ. 2564)
// ที่มา: โปรแกรม "กชกร Version 2" (kotchakorn_v2_20210404.xlsx, นพ.ชลทิศ อุไรฤกษ์กุล) ชีต age / ht
// - hfa / wfa: อายุ (เดือนเต็ม) 60–228 · แถว 72–228 = ตารางกรมอนามัย · 61–71 = ค่าคำนวณเชื่อมช่วงตามโปรแกรมกชกร · 60 = WHO 2006
//   แต่ละแถว [-2SD, -1.5SD, มัธยฐาน, +1.5SD, +2SD]
// - wfh: ส่วนสูง (ซม.) ชาย 120–185 · หญิง 120–175 ทีละ 1 ซม. — แถว [-2SD, -1.5SD, มัธยฐาน, +1.5SD, +2SD, +3SD]
//   (ต่ำกว่า 120 ซม. growth.js ใช้ WHO 2006)
// สร้างจากไฟล์ด้วยสคริปต์ ห้ามแก้ตัวเลขด้วยมือ
export default {''']
for k,label in (('wfh','น้ำหนักตามเกณฑ์ส่วนสูง'),('hfa','ส่วนสูงตามเกณฑ์อายุ'),('wfa','น้ำหนักตามเกณฑ์อายุ')):
    js.append(f'  // {label}\n  {k}: {{\n'+',\n'.join(f'  "{s}": '+tab(out[k][s]) for s in ('ช','ญ'))+'\n  },')
js.append('};\n')
open(sys.argv[2],'w').write('\n'.join(js))
